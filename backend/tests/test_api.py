from fastapi.testclient import TestClient

from app.main import create_app
from app.auth import hash_password
from app.models import Base, User


def make_client(tmp_path):
    app = create_app(
        f"sqlite:///{(tmp_path / 'test.db').as_posix()}",
        seed=False,
        attachment_dir=tmp_path / "files",
        collection_dir=tmp_path / "inbox",
    )
    Base.metadata.create_all(app.state.engine)
    with app.state.session_factory() as db:
        db.add(User(username="investigator", password_hash=hash_password("test-password-123"), role="investigator"))
        db.commit()
    client = TestClient(app)
    with client:
        login = client.post("/api/auth/login", json={"username": "investigator", "password": "test-password-123"})
        client.headers["X-CSRF-Token"] = login.json()["csrf"]
    return client


def test_health_and_meta(tmp_path):
    with make_client(tmp_path) as client:
        health = client.get("/api/health")
        assert health.status_code == 200
        assert health.json()["service"] == "zanaq-forensic-smart"
        meta = client.get("/api/meta")
        assert "Employee Fraud" in meta.json()["case_types"]
        assert meta.json()["transitions"]["new"] == ["in_review", "closed"]


def test_case_workflow_sorts_by_score_and_records_notes(tmp_path):
    with make_client(tmp_path) as client:
        low = client.post(
            "/api/cases",
            json={"title": "Low priority review", "case_type": "Other", "score": 15, "actor": "Intake desk"},
        )
        high = client.post(
            "/api/cases",
            json={
                "title": "High priority leakage review",
                "case_type": "Information Leakage",
                "score": 80,
                "actor": "Intake desk",
            },
        )
        assert low.status_code == 201
        assert high.status_code == 201
        high_id = high.json()["id"]

        listed = client.get("/api/cases")
        assert [row["title"] for row in listed.json()] == [
            "High priority leakage review",
            "Low priority review",
        ]

        blocked = client.post(
            f"/api/cases/{high_id}/transition",
            json={"status": "escalated", "actor": "Fraud investigator"},
        )
        assert blocked.status_code == 409

        moved = client.post(
            f"/api/cases/{high_id}/transition",
            json={"status": "in_review", "actor": "Fraud investigator", "note": "Queued for today"},
        )
        assert moved.status_code == 200
        assert moved.json()["status"] == "in_review"

        noted = client.post(
            f"/api/cases/{high_id}/notes",
            json={"author": "Fraud investigator", "body": "Compared the export with last week's report."},
        )
        assert noted.status_code == 201
        assert noted.json()["notes"][-1]["body"] == "Compared the export with last week's report."
        assert any(item["action"] == "note_added" for item in noted.json()["activities"])


def test_alert_link_raises_case_score_and_dismiss_blocks_relink(tmp_path):
    with make_client(tmp_path) as client:
        created = client.post(
            "/api/cases",
            json={"title": "Shared login follow-up", "case_type": "Identity Theft", "score": 10},
        )
        case_id = created.json()["id"]
        alert = client.post(
            "/api/alerts",
            json={
                "title": "Unusual shared-login export",
                "score": 77,
                "entity_type": "user",
                "entity_ref": "OPS-SHARED-2",
                "channel": "Back office",
                "case_id": case_id,
            },
        )
        assert alert.status_code == 201
        detail = client.get(f"/api/cases/{case_id}")
        assert detail.json()["score"] == 77
        assert detail.json()["alert_count"] == 1

        filtered = client.get("/api/cases", params={"min_score": 70})
        assert [row["id"] for row in filtered.json()] == [case_id]

        alert_id = alert.json()["id"]
        dismissed = client.patch(f"/api/alerts/{alert_id}", json={"status": "dismissed", "actor": "Duty supervisor"})
        assert dismissed.status_code == 200
        assert dismissed.json()["case_id"] is None

        relink = client.patch(f"/api/alerts/{alert_id}", json={"case_id": case_id})
        assert relink.status_code == 409


def test_attachments_and_report_exports(tmp_path):
    from io import BytesIO

    from openpyxl import load_workbook

    with make_client(tmp_path) as client:
        created = client.post(
            "/api/cases",
            json={"title": "Export sample case", "case_type": "Other", "score": 30, "summary": "Ready for export"},
        )
        assert created.status_code == 201
        case_id = created.json()["id"]
        client.post(
            "/api/cases",
            json={"title": "Below the export cutoff", "case_type": "Other", "score": 5},
        )

        uploaded = client.post(
            f"/api/cases/{case_id}/attachments",
            files={"file": ("..\\roster.txt", b"Roster comparison", "text/plain")},
        )
        assert uploaded.status_code == 201
        assert uploaded.json()["original_name"] == "roster.txt"
        attachment_id = uploaded.json()["id"]

        blocked = client.post(
            f"/api/cases/{case_id}/attachments",
            files={"file": ("run.exe", b"nope", "application/octet-stream")},
        )
        assert blocked.status_code == 400

        client.app.state.max_attachment_bytes = 8
        oversized = client.post(
            f"/api/cases/{case_id}/attachments",
            files={"file": ("note.txt", b"0123456789", "text/plain")},
        )
        assert oversized.status_code == 413

        downloaded = client.get(f"/api/cases/{case_id}/attachments/{attachment_id}")
        assert downloaded.status_code == 200
        assert downloaded.content == b"Roster comparison"

        detail = client.get(f"/api/cases/{case_id}")
        assert detail.json()["attachments"][0]["original_name"] == "roster.txt"
        assert any(item["action"] == "attachment_added" for item in detail.json()["activities"])

        removed = client.delete(f"/api/cases/{case_id}/attachments/{attachment_id}")
        assert removed.status_code == 204
        assert client.get(f"/api/cases/{case_id}/attachments/{attachment_id}").status_code == 404

        csv_export = client.get("/api/reports/cases.csv", params={"min_score": 20})
        assert csv_export.status_code == 200
        assert "Export sample case" in csv_export.text
        assert "Below the export cutoff" not in csv_export.text

        workbook = load_workbook(BytesIO(client.get("/api/reports/cases.xlsx").content))
        titles = [row[1] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
        assert "Export sample case" in titles

        pdf = client.get(f"/api/reports/cases/{case_id}.pdf")
        assert pdf.status_code == 200
        assert pdf.content.startswith(b"%PDF")
        assert b"Export sample case" in pdf.content
